"""Shared logic for the translation table (build_tm.py) and putting translations back (apply_translations.py).

Core text: each text line is split into  prefix + core + suffix.
  prefix = leading {tags} and [notes] (e.g. "{AngryA}[pickup][ALT]"), copied back untouched;
  core   = what the translator translates (mid-line tags stay inside it);
  suffix = trailing spaces/tabs.
  A line that is only brackets ("[Ignore]", "[sighs]") keeps its brackets in the core.

Ukrainian gender: a core is split into separate table rows by speaker gender when it has first-person
past tense / "I'm sure"-type words, and by listener gender when it has "you" + past tense / adjective.
Speaker gender comes from reference/speaker_genders.xlsx; the listener is guessed (see guess_listeners).
"""
import collections
import hashlib
import os
import sys
import re

NOTE = re.compile(r'^\(\w[\w .]*, \w{3} \w{3}')          # writers' timestamped notes
MARKUP = re.compile(r'\{[^{}]*\}|\^[^^]*\^|<<[^<>]*>>|<[^<>]+>')
PREFIX = re.compile(r'^(?:\s*(?:\{[^{}]*\}|\[[^\[\]]*\]))+\s*')
TAGS_ONLY_PREFIX = re.compile(r'^(?:\s*\{[^{}]*\})+\s*')
DEV_SPEAKER = re.compile(r'^(DIR|STUB|NOTES?|QUESTION|GAMEPLAY|CHANGED?|WRITING|VFX|QTE|FILE|HTTPS?|REF|TESTDIR)')
LINE_SEP = ' ¶ '                                         # joins lines of a multi-line record in the table

_IRREG = ('was|were|did|saw|made|had|went|got|told|said|thought|knew|found|left|took|gave|came|felt|kept|lost|'
          'met|heard|meant|brought|bought|caught|fought|forgot|forgave|ran|sat|stood|understood|wrote|became|'
          'began|broke|chose|drove|fell|flew|grew|hid|hit|hurt|let|put|read|rode|sent|shot|slept|spoke|spent|'
          'stole|swore|taught|threw|tried|won|woke|could|would')
_ADJ = ('sure|glad|sorry|ready|tired|afraid|happy|alone|wrong|right|done|busy|serious|grateful|curious|aware|'
        'late|free|alive|dead|crazy|drunk|innocent|guilty|proud|scared|worried|surprised|impressed|interested|'
        'disappointed|married|responsible|calm|safe|sick|honest|able|certain|positive|convinced|lucky|young|'
        'old|good|fine|okay|alright|new|mad|angry|upset|lonely|jealous|stupid|smart|strong|weak|hungry|awake')
_ADV = r'(?:just |never |already |really |always |only |even |also |not |so |very |too |quite |a little |still |all )*'
SPEAKER_GENDERED = re.compile(
    rf"\bI(?:’d|'d|’ve|'ve| had| have| would have)?\s+{_ADV}(?:\w+ed|{_IRREG})\b"
    rf"|\bI(?:’m|'m| am| was)\s+{_ADV}(?:{_ADJ}|\w+ed)\b|\bmyself\b", re.I)
LISTENER_GENDERED = re.compile(
    rf"\byou(?:’d|'d|’ve|'ve| had| have| would have)?\s+{_ADV}(?:\w+ed|{_IRREG})\b"
    rf"|\byou(?:’re|'re| are| were)\s+{_ADV}(?:{_ADJ}|\w+ed)\b|\b(?:did|were|have|had) you\b|\byourself\b", re.I)

ALIAS = {'BATMANBRUCE': 'BATMAN', 'DENT': 'HARVEY', 'HARBEY': 'HARVEY', 'RENEE': 'MONTOYA', 'JOE CHILL': 'JOECHILL'}
VOCATIVES = {
    'BRUCE': r'Bruce|Mr\. Wayne|Master Bruce|Master Wayne|Brucie|Wayne', 'BATMAN': r'Batman|Bats',
    'ALFRED': r'Alfred|Pennyworth', 'HARVEY': r'Harvey|Harv|Mr\. Dent|Dent|Mayor Dent',
    'SELINA': r'Selina|Ms\. Kyle|Miss Kyle|Kyle', 'CATWOMAN': r'Catwoman',
    'GORDON': r'Gordon|Jim|Lieutenant|Commissioner', 'OSWALD': r'Oswald|Oz|Ozzy|Cobblepot', 'PENGUIN': r'Penguin',
    'VICKI': r'Vicki|Ms\. Vale|Miss Vale', 'LUCIUS': r'Lucius|Mr\. Fox', 'FALCONE': r'Falcone|Carmine',
    'MONTOYA': r'Montoya|Renee', 'HILL': r'Mayor Hill|Hill|Mr\. Mayor', 'JOKER': r'John|Doe',
    'LELAND': r'Dr\. Leland|Leland|Doctor', 'LADYARKHAM': r'Lady Arkham', 'ZSASZ': r'Zsasz|Victor',
    'REGINA': r'Regina|Mrs\. Zellerbach'}
_VOC_RE = {k: re.compile(rf'(?:^|[,!?.]\s*|--\s*)(?:{v})(?:\s*[,!?.]|\s*--|\s*$)') for k, v in VOCATIVES.items()}
_TO_RE = re.compile(r'\[(?:[^\]]*;\s*)?to ([A-Za-z .’\-]+?)\s*[;\]]', re.I)


def check_root(files, root):
    """Stop early if EN_ROOT is not the folder that directly contains the BM_* / BAT_* folders."""
    if not files or not all(f.startswith(('BM_', 'BAT_')) for f in files):
        sys.exit(f'{root} must be the folder that directly contains the BM_* and BAT_* folders '
                 f'(here: original_data). Found e.g.: {files[:1]}')


def load_genders(path):
    """speaker_id -> 'M' / 'F' / '?'  from speaker_genders.xlsx (or .tsv)."""
    rows = []
    if path.lower().endswith('.xlsx'):
        _openpyxl()
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        next(it)
        rows = [[('' if v is None else str(v)) for v in r[:2]] for r in it]
        wb.close()
    else:
        with open(path, encoding='utf-8-sig') as fh:
            next(fh)
            rows = [line.rstrip('\r\n').split('\t') for line in fh]
    return {r[0].strip(): (r[1].strip().upper() or '?') for r in rows if len(r) >= 2 and r[0].strip()}


def norm_speaker(s):
    return ALIAS.get(s, s)


def is_dev_record(rec):
    return bool(DEV_SPEAKER.match(rec['spk'])) or any(NOTE.match(l) for l in rec['body'][:1])


TRAILING_ALT = re.compile(r'\s*\[alt\b[^\[\]]*\]$', re.I)


def split_line(line):
    """-> (prefix, core, suffix)."""
    stripped = line.rstrip(' \t')
    suffix = line[len(stripped):]
    m = TRAILING_ALT.search(stripped)
    if m and stripped[:m.start()].strip():
        suffix = stripped[m.start():] + suffix
        stripped = stripped[:m.start()]
    m = PREFIX.match(stripped)
    if m and stripped[m.end():].strip():
        return stripped[:m.end()], stripped[m.end():], suffix
    m = TAGS_ONLY_PREFIX.match(stripped)           # only brackets left: keep them in the core
    if m and stripped[m.end():].strip():
        return stripped[:m.end()], stripped[m.end():], suffix
    return '', stripped, suffix


def has_letters(text):
    return any(c.isalpha() for c in MARKUP.sub('', text))


def spoken(text):
    return re.sub(r'\[[^\]]*\]', ' ', MARKUP.sub('', text))


def guess_listeners(recs):
    """recs sorted by id. Returns {id: (listener, how)} for records with a speaker (how = note/name/nearby).
    Records without a speaker (choices) are spoken by Bruce/Batman to the nearest other speaker."""
    seq = [(i, norm_speaker(r['spk'])) for i, r in enumerate(recs)
           if r['spk'] and not is_dev_record(r) and any(l.strip() for l in r['body'])]
    pos = {i: k for k, (i, _) in enumerate(seq)}
    out = {}
    for i, r in enumerate(recs):
        if is_dev_record(r):
            continue
        me = norm_speaker(r['spk']) if r['spk'] else 'BRUCE'
        text = MARKUP.sub('', ' '.join(r['body']))
        target, how = None, None
        m = _TO_RE.search(text)
        if m:
            name = re.sub(r'[^A-Z]', '', m.group(1).upper())
            for k, rx in VOCATIVES.items():
                alts = [re.sub(r'[^A-Z]', '', a.upper()) for a in rx.replace('\\.', '').split('|')]
                if name == k or name in alts:
                    target, how = k, 'note'
                    break
        if not target:
            sp = spoken(text)
            for k, rx in _VOC_RE.items():
                if k != me and rx.search(sp):
                    target, how = k, 'name'
                    break
        if not target:
            if i in pos:
                k0 = pos[i]
                near = [s for _, s in seq[max(0, k0 - 3):k0 + 4] if s != me]
            else:  # choice: look at speakers around it
                around = [s for j, s in seq if abs(j - i) <= 4 and s not in ('BRUCE', 'BATMAN')]
                near = around
            if near:
                target, how = collections.Counter(near).most_common(1)[0][0], 'nearby'
        out[r['id']] = (target, how)
    return out


def record_key(cores, speaker_gender, listener_gender):
    """Stable id of a table row. Gender parts are added only when the text needs them."""
    joined = '\n'.join(cores)
    sp = spoken(joined)
    sg = speaker_gender if SPEAKER_GENDERED.search(sp) else ''
    lg = listener_gender if LISTENER_GENDERED.search(sp) else ''
    h = hashlib.sha1(joined.encode('utf-8')).hexdigest()[:10]
    return h + (f'-S{sg}' if sg else '') + (f'-L{lg}' if lg else ''), sg, lg


def analyze_records(recs, genders):
    """recs: records of one file. Returns {id: info} for every record that belongs in the table.
    info = dict(cores, key, sg, lg, speaker, listener, how). Missing id = stays English (notes, no letters,
    subtitles that are only [sighs])."""
    by_id = sorted(recs, key=lambda r: r['id'])
    listeners = guess_listeners(by_id)
    out = {}
    for r in by_id:
        if is_dev_record(r):
            continue
        cores = [split_line(l)[1] for l in r['body']]
        if not has_letters(' '.join(cores)):
            continue
        if r['spk'] and not has_letters(re.sub(r'\[[^\[\]]*\]', '', ' '.join(cores))):
            continue
        spk = norm_speaker(r['spk']) if r['spk'] else 'BRUCE'
        lst, how = listeners.get(r['id'], (None, None))
        sg_all = genders.get(r['spk'], genders.get(spk, '?'))
        lg_all = genders.get(lst, '?') if lst else '?'
        key, sg, lg = record_key(cores, sg_all, lg_all)
        out[r['id']] = {'cores': cores, 'key': key, 'sg': sg, 'lg': lg,
                        'speaker': r['spk'] or 'choice/UI', 'listener': lst, 'how': how}
    return out


def split_translation(cell, n_lines):
    """Table cell -> list of n_lines core lines, or None if the ¶ count is wrong."""
    cell = unescape_cell(cell)
    if n_lines == 1:
        return [cell]
    parts = cell.split(LINE_SEP)               # exact " ¶ " keeps spaces inside lines
    if len(parts) != n_lines:
        parts = re.split(r' ?¶ ?', cell)       # editor trimmed a space somewhere
    return parts if len(parts) == n_lines else None


def loose_text(s):
    """Lower-case letters/digits only - for spotting near-identical rows."""
    s = re.sub(r'\{[^}]*\}|\[[^\]]*\]|\^[^^]*\^|<[^>]*>', '', s).replace('’', "'").lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


# --- translation table files: .xlsx (for people) or .tsv (plain text) ------------------------------------
# (internal name, column title in Excel, width)
TABLE_COLUMNS = [
    ('row', '#', 7),
    ('english', 'English', 60),
    ('ukrainian', 'Ukrainian  ✎ (type here)', 60),
    ('speaker_gender', 'Speaker gender', 9),
    ('listener_gender', 'Listener gender', 9),
    ('speakers', 'Who speaks', 18),
    ('listeners', 'To whom (guess)', 18),
    ('count', 'Used (times)', 8),
    ('first_seen', 'Where (file : line)', 45),
    ('notes', 'Notes', 40),
    ('key', 'key (do not edit)', 16),
]
HELP = [
    'How to use this file',
    '',
    'Type the Ukrainian translation into the yellow column "Ukrainian". Do not change any other column.',
    'One row = one text. It is used in every place the game shows it ("Used (times)" says how many).',
    'Empty Ukrainian cell = the game keeps the English text.',
    '',
    'Columns',
    'English - the text to translate. Tags at the start of a line ({AngryA}, [pickup], [ALT]) are already removed;',
    '    the script puts them back. Tags still inside the text ({SmugA}, ^font..^, <icon.d3dtx:-15>) must be copied',
    '    unchanged into the translation.',
    'Speaker gender - filled only if the Ukrainian words depend on who speaks: M = male (бачив, певен),',
    '    F = female (бачила, певна), ? = unknown, check the scene. Empty = gender does not matter here.',
    'Listener gender - same, for the person spoken to ("Did you see?" -> бачив / бачила).',
    'Who speaks / To whom - characters. "choice/UI" = a choice option or menu text (Bruce says choices).',
    '    "To whom" is a guess from the scene.',
    'Where - first place this text appears. Open reference/scenes/<that file> to read the whole conversation.',
    'Notes - "keep N ¶": the text has several lines joined with " ¶ " - keep the same number of ¶.',
    '    "near-copy of <key>": almost the same as another row - reuse that translation.',
    '    "listener guessed - check scene": the gender of the listener is a guess.',
    '    "credits: people’s names stay in Latin".',
    'key - technical id that links the row to the game files. Never edit it, never delete rows.',
    '',
    'Rules',
    '- Keep "...", "--" and the ending punctuation the same as in the English text.',
    '- Bracket-only choices like [Ignore] are visible in game: translate inside the brackets, keep the brackets.',
    '- Use names exactly as in reference/glossary.md, and ти/ви as in reference/speakers.md.',
    '- Turn off AutoCorrect if it changes "..." into "…" or "--" into a dash.',
    '- The Ukrainian column is formatted as Text, so lines starting with -- or = are safe to type.',
    '- You can sort and filter, but save as .xlsx (not .csv).',
]


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        sys.exit('Excel files need the openpyxl package. Install it once with:  python -m pip install openpyxl')


def write_table(path, rows):
    """rows: list of dicts with internal column names."""
    names = [c[0] for c in TABLE_COLUMNS]
    if path.lower().endswith('.xlsx'):
        _openpyxl()
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        wb = Workbook()
        ws = wb.active
        ws.title = 'Translate'
        ws.append([c[1] for c in TABLE_COLUMNS])
        wrap = Alignment(wrap_text=True, vertical='top')
        top = Alignment(vertical='top')
        yellow = PatternFill('solid', fgColor='FFF2B3')
        header_fill = PatternFill('solid', fgColor='E7E6E6')
        grey = Font(color='999999', size=8)
        line = Side(style='thin', color='BFBFBF')
        box = Border(left=line, right=line, top=line, bottom=line)  # fills hide Excel's gridlines
        for n, row in enumerate(rows, 1):
            ws.append([n if c == 'row' else row.get(c, '') for c in names])
        for idx, (name, title, width) in enumerate(TABLE_COLUMNS, 1):
            letter = ws.cell(row=1, column=idx).column_letter
            ws.column_dimensions[letter].width = width
            head = ws.cell(row=1, column=idx)
            head.font = Font(bold=True)
            head.fill = header_fill
            head.border = box
            head.alignment = wrap
        ucol = names.index('ukrainian') + 1
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.border = box
                cell.alignment = wrap if cell.column in (2, ucol, len(names) - 1) else top
                if cell.column != 1 and cell.column != names.index('count') + 1:
                    cell.data_type = 's'           # never a formula, even if text starts with = or --
            u = r[ucol - 1]
            u.fill = yellow
            u.number_format = '@'
            r[-1].font = grey
        ws.cell(row=1, column=ucol).fill = yellow
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        hs = wb.create_sheet('Help')
        hs.column_dimensions['A'].width = 120
        for line in HELP:
            hs.append([line])
        hs['A1'].font = Font(bold=True, size=14)
        wb.save(path)
    else:
        with open(path, 'w', encoding='utf-8-sig', newline='\r\n') as fh:
            fh.write('\t'.join(n for n in names if n != 'row') + '\n')
            for row in rows:
                fh.write('\t'.join(str(row.get(n, '')) for n in names if n != 'row') + '\n')


def read_table(path):
    """-> list of row dicts with internal column names (from .xlsx or .tsv)."""
    if path.lower().endswith('.xlsx'):
        _openpyxl()
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        ws = wb['Translate'] if 'Translate' in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        titles = [str(t or '') for t in next(it)]
        by_title = {c[1]: c[0] for c in TABLE_COLUMNS}
        cols = [by_title.get(t, t) for t in titles]
        rows = []
        for vals in it:
            row = {c: ('' if v is None else str(v)) for c, v in zip(cols, vals)}
            if row.get('key'):
                rows.append(row)
        wb.close()
        return rows
    with open(path, encoding='utf-8-sig') as fh:
        header = fh.readline().rstrip('\r\n').split('\t')
        return [dict(zip(header, line.rstrip('\r\n').split('\t'))) for line in fh if line.strip()]


# --- unique-text folders (unique_original_data / unique_localized_data) ----------------------------------
# unique_original_data is built from the English game files ONLY (no genders, no guesses): every unique English
# text once, in the file named after the game file where it first appears:
#   #### <text id> | <who speaks, from the header lines> → <spoken to, guessed from the scene> | x<times used>
# ("spoken to" comes from [to X] notes, names used in direct address and neighbouring speakers - all from the
#  English files, never from the genders file)
#   <text lines>
#   <empty line>
# unique_localized_data has the same entries with the Ukrainian from the table (untranslated = English). Only when
# one English text needs several Ukrainian versions (male/female forms) does the entry appear once per version,
# with the table key and a label, e.g. "#### f156de4535-LF | BRUCE | x4 | listener: female".

def file_order(rel):
    """Default mode first (episode 101 -> 105, menus), then Noir; used for table rows and unique files."""
    folder = rel.split('/')[0]
    mode = 0 if folder.startswith('BM_') else 1
    m = re.search(r'(10[1-5])', folder)
    grp = m.group(1) if m else {'Menu': '200', 'Boot': '300', 'Project': '400'}.get(
        next((g for g in ('Menu', 'Boot', 'Project') if g in folder), ''), '900')
    return mode, grp, rel


def english_unique(root):
    """Every unique English text of the game files, in table order: {text id: entry}. Uses no genders."""
    from landb_txt import read_records, txt_files
    out = {}
    for rel in sorted(txt_files(root), key=file_order):
        recs = read_records(os.path.join(root, rel))[0]
        info = analyze_records(recs, {})
        for r in sorted(recs, key=lambda x: x['id']):
            if r['id'] not in info:
                continue
            tid = info[r['id']]['key'].split('-')[0]
            e = out.get(tid)
            if e is None:
                e = out[tid] = {'cores': info[r['id']]['cores'], 'file': rel,
                                'speakers': collections.Counter(), 'listeners': collections.Counter(), 'count': 0}
            e['count'] += 1
            e['speakers'][r['spk'] or 'choice/UI'] += 1
            if info[r['id']]['listener']:
                e['listeners'][info[r['id']]['listener']] += 1
    return out


def _head(tid, e):
    who = ', '.join(s for s, _ in e['speakers'].most_common(4))
    to = ', '.join(s for s, _ in e['listeners'].most_common(3))
    return f"#### {tid} | {who}{' → ' + to if to else ''} | x{e['count']}"


GENDER_WORD = {'M': 'male', 'F': 'female', '?': 'unknown'}


def unique_files(root, rows=None):
    """-> ({relative path: [lines]}, rows with a wrong ¶ count). rows=None: English folder.
    rows = table rows: Ukrainian folder."""
    eng = english_unique(root)
    variants = collections.defaultdict(list)
    for row in rows or []:
        variants[row['key'].split('-')[0]].append(row)
    files = collections.defaultdict(list)
    bad = 0
    for tid, e in eng.items():
        if rows is None or len(variants.get(tid, [])) <= 1:
            text = e['cores']
            row = variants[tid][0] if rows is not None and variants.get(tid) else None
            if row and row.get('ukrainian'):
                tr = split_translation(row['ukrainian'], len(e['cores']))
                if tr is None:
                    bad += 1
                else:
                    text = tr
            files[e['file']].extend([_head(tid, e)] + text + [''])
            continue
        for row in variants[tid]:                          # several Ukrainian versions of one English text
            label = ', '.join(f'{who}: {GENDER_WORD.get(g, g)}' for who, g in
                              (('speaker', row.get('speaker_gender')), ('listener', row.get('listener_gender'))) if g)
            text = e['cores']
            if row.get('ukrainian'):
                tr = split_translation(row['ukrainian'], len(e['cores']))
                if tr is None:
                    bad += 1
                else:
                    text = tr
            files[e['file']].extend([_head(row['key'], e) + f' | {label}'] + text + [''])
    return files, bad


def _unique_folder_files(d):
    """All files in d. Exits if d holds anything that is not a unique-text file (protects other folders)."""
    from landb_txt import read_raw
    found = []
    for dp, _, fn in os.walk(d):
        for f in fn:
            path = os.path.join(dp, f)
            ok = f.endswith('.txt') and read_raw(path)[:20].decode('utf-8', 'replace').lstrip('﻿').startswith('#### ')
            if not ok:
                sys.exit(f'{d} contains {os.path.relpath(path, d)}, which is not a unique-text file - '
                         f'refusing to touch this folder. Check the folder name.')
            found.append(os.path.relpath(path, d).replace(os.sep, '/'))
    return found


def write_unique(files, d):
    """Rewrite folder d with files from unique_files(). Only unique-text files are ever deleted."""
    from landb_txt import write_lines
    for rel in _unique_folder_files(d):
        if rel not in files:
            os.remove(os.path.join(d, rel))
    for rel, lines in files.items():
        write_lines(os.path.join(d, rel), lines)
    return len(files)


def compare_unique(root, d):
    """Is folder d exactly the English unique folder of root? -> list of problems (empty = yes)."""
    from landb_txt import read_raw, split_lines
    files, _ = unique_files(root)
    have = set(_unique_folder_files(d)) if os.path.isdir(d) else set()
    problems = [f'missing: {rel}' for rel in sorted(set(files) - have)]
    problems += [f'not from the English files: {rel}' for rel in sorted(have - set(files))]
    for rel in sorted(set(files) & have):
        if split_lines(read_raw(os.path.join(d, rel)))[0] != files[rel]:
            problems.append(f'differs: {rel}')
    return problems


def escape_cell(s):
    return s.replace('\\', '\\\\').replace('\t', '\\t')


def unescape_cell(s):
    return re.sub(r'\\(\\|t)', lambda m: '\t' if m.group(1) == 't' else '\\', s)
