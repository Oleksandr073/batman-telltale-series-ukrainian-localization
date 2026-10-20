"""End-to-end self-test of the translation flow, on random parts of the game, in a temporary copy.

Usage (from the repo root):
  python .claude/skills/telltale-batman-localization/scripts/selftest.py [--seed N] [--scenes K] [--keep]

Nothing in the repo is changed: data/original_data, data/unique_original_data, speaker_genders.xlsx and
ai_context/play_order.txt are copied to a temporary folder, and every script runs there.

What it does:
  1. builds a fresh table (build_tm.py) and fills random scenes from random episodes, a menu file and Noir-only texts
     with fake "Ukrainian" (English letters swapped for Cyrillic, all tags kept). Male/female versions get different
     text; some multi-line texts use Excel line breaks (Alt+Enter) instead of " ¶ ".
  2. runs the import (apply_translations.py) and then checks EVERY record of the 222 game files independently:
     translated records = the fake text with the original leading tags / [ALT] / trailing spaces restored;
     every other record byte-identical to English; headers, line counts, BOM and CRLF unchanged.
  3. checks the unique folders, the grey reference rows, that a second import gives identical files, that a rebuild
     keeps translations and notes, and that a gender change keeps translations (marked "check gender forms").
  4. makes deliberate mistakes (line break in a one-line text, typing in a reference row, a deleted tag) and checks
     that each one is reported.
Exit code 0 = all checks passed.
"""
import argparse
import hashlib
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from landb_txt import BOM, read_raw, read_records, split_lines, txt_files  # noqa: E402
from tm_lib import (LINE_SEP, TABLE_COLUMNS, analyze_records, escape_cell, load_genders, split_line,  # noqa: E402
                    unescape_cell)

RESULTS = []
TITLE = {c[0]: c[1] for c in TABLE_COLUMNS}

# ---------------------------------------------------------------------------------------------------------
# fake Ukrainian: English letters -> Cyrillic, using all 33 Ukrainian letters so the game font is fully tested
# (ч ш ґ ю я є ї ь come from letter pairs, e.g. night -> ніґт, cheer -> чєр, all -> аль)
PAIRS = {'ch': 'ч', 'sh': 'ш', 'gh': 'ґ', 'oo': 'ю', 'ea': 'я', 'ee': 'є', 'ie': 'ї', 'll': 'ль'}
SINGLE = dict(zip('abcdefghijklmnopqrstuvwxyz', ['а', 'б', 'ц', 'д', 'е', 'ф', 'г', 'х', 'і', 'й', 'к', 'л', 'м', 'н',
                                                 'о', 'п', 'щ', 'р', 'с', 'т', 'у', 'в', 'ш', 'ж', 'и', 'з']))
FAKE_LETTERS = re.compile('|'.join(PAIRS) + '|[a-z]', re.I)


def _swap(m):
    s = m.group(0)
    c = PAIRS.get(s.lower()) or SINGLE[s.lower()]
    return c.upper() if s[0].isupper() else c


KEEP = re.compile(r'(\{[^{}]*\}|\^[^^]*\^|<<[^<>]*>>|<[^<>]+>| ¶ |\\t|\[alt\b[^\[\]]*\])', re.I)


def fake(cell):
    """English table cell -> fake Ukrainian with every tag, ¶ and \\t kept."""
    return ''.join(p if KEEP.fullmatch(p) else FAKE_LETTERS.sub(_swap, p) for p in KEEP.split(cell))


def check(name, ok, detail=''):
    RESULTS.append((name, bool(ok), detail))
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f'  - {detail}' if detail else ''))
    return ok


def run(args, cwd):
    p = subprocess.run([sys.executable] + args, cwd=cwd, capture_output=True, text=True, encoding='utf-8',
                       errors='replace')
    return p.returncode, p.stdout + p.stderr


def tree_hash(root):
    h = hashlib.sha1()
    for rel in txt_files(root):
        h.update(rel.encode())
        h.update(read_raw(os.path.join(root, rel)))
    return h.hexdigest()


# --- table helpers (openpyxl, like a person editing in Excel) ----------------------------------------------
def load_sheet(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb['Translate']
    col = {name: [c.value for c in ws[1]].index(TITLE[name]) + 1 for name in TITLE}
    return wb, ws, col


def table_rows(ws, col):
    """-> list of dicts: row, kind, english, key, ukrainian, notes, lines, speaker_gender, listener_gender,
    speakers, first_seen"""
    out = []
    for r in range(2, ws.max_row + 1):
        g = lambda n: ws.cell(row=r, column=col[n]).value  # noqa: E731
        key, notes = g('key'), str(g('notes') or '')
        if key:
            kind = 'text'
        elif notes.startswith('see row #'):
            kind = 'ref'
        elif ws.row_dimensions[r].outlineLevel == 0:
            kind = 'section'
        else:
            kind = 'scene'
        out.append({'row': r, 'kind': kind, 'english': g('english') or '', 'key': key or '',
                    'ukrainian': g('ukrainian') or '', 'notes': notes, 'lines': g('lines'),
                    'speaker_gender': g('speaker_gender') or '', 'listener_gender': g('listener_gender') or '',
                    'speakers': g('speakers') or '', 'first_seen': g('first_seen') or ''})
    return out


# --- independent expectation for every record ---------------------------------------------------------------
def expected_lines(body, cell):
    """What the game file must contain for a record whose table cell is `cell` (as typed in Excel)."""
    n = len(body)
    text = unescape_cell(cell)
    parts = text.split('\n') if '\n' in text else (text.split(LINE_SEP) if n > 1 else [text])
    assert len(parts) == n, (body, cell)
    out = []
    for line, part in zip(body, parts):
        prefix, _, suffix = split_line(line)
        out.append('' if line == '' else prefix + part + suffix)
    return out


def verify_game_files(en_root, loc_root, genders, filled):
    """Every record of every file: translated as expected, or byte-identical English."""
    problems, n_tr, n_same = [], 0, 0
    files = txt_files(en_root)
    for rel in files:
        en_raw, loc_raw = read_raw(os.path.join(en_root, rel)), read_raw(os.path.join(loc_root, rel))
        if not loc_raw.startswith(BOM) or loc_raw.replace(b'\r\n', b'').count(b'\n') or not loc_raw.endswith(b'\r\n'):
            problems.append(f'{rel}: BOM/CRLF broken')
        en_lines, loc_lines = split_lines(en_raw)[0], split_lines(loc_raw)[0]
        if len(en_lines) != len(loc_lines):
            problems.append(f'{rel}: {len(loc_lines)} lines, English has {len(en_lines)}')
            continue
        en_recs = read_records(os.path.join(en_root, rel))[0]
        loc_recs = read_records(os.path.join(loc_root, rel))[0]
        info = analyze_records(en_recs, genders)
        for e, l in zip(en_recs, loc_recs):
            if e['hdr'] != l['hdr']:
                problems.append(f"{rel}:{e['idx'] + 1}: header changed")
                continue
            i = info.get(e['id'])
            cell = filled.get(i['key']) if i else None
            want = expected_lines(e['body'], cell) if cell else e['body']
            if l['body'] != want:
                problems.append(f"{rel}:{e['idx'] + 1}: expected {want!r:.80} found {l['body']!r:.80}")
            elif cell:
                n_tr += 1
            else:
                n_same += 1
    return problems, n_tr, n_same, len(files)


# -------------------------------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--seed', type=int, default=None, help='random seed (default: new each run, printed)')
    ap.add_argument('--scenes', type=int, default=4, help='how many random scenes to translate (default 4)')
    ap.add_argument('--keep', action='store_true', help='keep the temporary folder for inspection')
    a = ap.parse_args()
    seed = a.seed if a.seed is not None else random.randrange(1_000_000)
    rnd = random.Random(seed)
    repo = os.getcwd()
    for need in ('data/original_data', 'data/unique_original_data', 'speaker_genders.xlsx', 'ai_context/play_order.txt'):
        if not os.path.exists(os.path.join(repo, need)):
            sys.exit(f'Run this from the repo root ({need} not found).')
    tmp = tempfile.mkdtemp(prefix='bm_selftest_')
    t0 = time.time()
    print(f'Self-test, seed {seed} (repeat with --seed {seed}). Working copy: {tmp}')
    try:
        for sub in ('data/original_data', 'data/unique_original_data'):
            shutil.copytree(os.path.join(repo, sub), os.path.join(tmp, sub))
        os.makedirs(os.path.join(tmp, 'ai_context'))
        shutil.copy(os.path.join(repo, 'ai_context/play_order.txt'), os.path.join(tmp, 'ai_context'))
        shutil.copy(os.path.join(repo, 'speaker_genders.xlsx'), tmp)
        S = lambda n: os.path.join(HERE, n)  # noqa: E731
        build = [S('build_tm.py'), 'data/original_data', 'speaker_genders.xlsx', 'translation_table.xlsx',
                 '--order', 'ai_context/play_order.txt']
        imp = [S('apply_translations.py'), 'data/original_data', 'speaker_genders.xlsx', 'translation_table.xlsx',
               'data/localized_data', '--unique', 'data/unique_original_data', 'data/unique_localized_data']
        uniq_before = tree_hash(os.path.join(tmp, 'data/unique_original_data'))
        genders = load_genders(os.path.join(tmp, 'speaker_genders.xlsx'))

        print('\n[1] Build a fresh table')
        code, out = run(build, tmp)
        check('table builds', code == 0, out.strip().splitlines()[-1] if out.strip() else '')

        print('\n[2] Fill random scenes with fake Ukrainian')
        wb, ws, col = load_sheet(os.path.join(tmp, 'translation_table.xlsx'))
        rows = table_rows(ws, col)
        scenes, cur = [], None
        section = ''
        for r in rows:
            if r['kind'] == 'section':
                section = r['english']
            elif r['kind'] == 'scene':
                cur = {'title': r['english'], 'section': section, 'rows': []}
                scenes.append(cur)
            elif cur is not None:
                cur['rows'].append(r)
        episodes = [s for s in scenes if s['section'].startswith('EPISODE') and any(x['kind'] == 'text' for x in s['rows'])]
        picked = rnd.sample(episodes, min(a.scenes, len(episodes)))
        picked += [rnd.choice([s for s in scenes if s['section'] == 'MENUS & UI'])]
        picked += [s for s in scenes if s['section'].startswith('SHADOWS MODE')][:1]
        texts = [r for r in rows if r['kind'] == 'text']
        # make sure the tricky kinds are covered even if the random scenes miss them
        extra = []
        for want in (lambda r: r['lines'] and r['lines'] > 1,                          # multi-line
                     lambda r: re.match(r'^[0-9a-f]+-[SL][MF]$', r['key']),            # male/female version
                     lambda r: 'keep markup inside' in r['notes'],                     # tags inside
                     lambda r: r['english'].startswith('[') and r['english'].endswith(']')):  # [Ignore] choice
            pool = [r for r in texts if want(r)]
            extra += rnd.sample(pool, min(3, len(pool)))
        variants = {}
        for r in texts:
            variants.setdefault(r['key'].split('-')[0], []).append(r)
        pair = [v for v in variants.values() if len(v) >= 2 and any(x['key'].endswith('F') for x in v)]
        if pair:
            extra += rnd.choice(pair)                                                  # both versions of one text
        to_fill = {r['key']: r for s in picked for r in s['rows'] if r['kind'] == 'text'}
        to_fill.update({r['key']: r for r in extra})
        filled, notes = {}, {}
        for key, r in to_fill.items():
            cell = fake(r['english'])
            if key.endswith('F'):
                cell += ' (ж)' if r['lines'] == 1 else ''
            if r['lines'] and r['lines'] > 1 and rnd.random() < 0.5:
                cell = cell.replace(LINE_SEP, '\n')                                   # Alt+Enter instead of ¶
            if key.endswith('F') and r['lines'] and r['lines'] > 1:
                cell = cell + ' (ж)'
            ws.cell(row=r['row'], column=col['ukrainian']).value = cell
            filled[key] = cell
        for r in rnd.sample(list(to_fill.values()), min(5, len(to_fill))):
            ws.cell(row=r['row'], column=col['translator_notes']).value = f'note for {r["key"]}'
            notes[r['key']] = f'note for {r["key"]}'
        wb.save(os.path.join(tmp, 'translation_table.xlsx'))
        print('  scenes: ' + ' | '.join(f"{s['section'].split(' —')[0]}: {s['title']}" for s in picked))
        check('random rows filled', len(filled) > 20, f'{len(filled)} texts, {len(notes)} translator notes')

        print('\n[3] Import and check every record of every game file')
        code, out = run(imp, tmp)
        check('import exit code 0', code == 0, re.findall(r'Validation: .*|Consistency: .*', out))
        problems, n_tr, n_same, n_files = verify_game_files(os.path.join(tmp, 'data/original_data'),
                                                            os.path.join(tmp, 'data/localized_data'), genders, filled)
        check('every record correct in all game files', not problems and n_tr > 0,
              f'{n_files} files, {n_tr} translated records, {n_same} unchanged English'
              + (f'; first problem: {problems[0]}' if problems else ''))
        bat_tr = 0
        for rel in txt_files(os.path.join(tmp, 'data/localized_data')):
            if rel.startswith('BAT_'):
                body = open(os.path.join(tmp, 'data/localized_data', rel), encoding='utf-8-sig').read()
                bat_tr += sum(1 for c in filled.values() if c.split('\n')[0].split(' ¶ ')[0][:25] in body)
        check('translations also reach Noir (BAT_*) files', bat_tr > 0, f'{bat_tr} hits')
        fem = [k for k in filled if k.endswith('F')]
        check('male/female versions routed separately', not fem or not problems, f'{len(fem)} female version(s) filled')
        check('data/unique_original_data untouched', tree_hash(os.path.join(tmp, 'data/unique_original_data')) == uniq_before)
        uk_text = ''.join(open(os.path.join(dp, f), encoding='utf-8-sig').read()
                          for dp, _, fn in os.walk(os.path.join(tmp, 'data/unique_localized_data')) for f in fn)
        missing = [k for k, c in filled.items() if unescape_cell(c.replace('\n', LINE_SEP)).split(LINE_SEP)[0] not in uk_text]
        check('each translation appears in data/unique_localized_data', not missing,
              f'missing {missing[:3]}' if missing else f'{len(filled)} found')
        check('no gender labels in unique folders', 'listener: ' not in uk_text and 'speaker: ' not in uk_text)

        refs = [r for r in rows if r['kind'] == 'ref']
        by_row = {r['row']: r for r in rows}
        bad_refs = [r['row'] for r in refs
                    if by_row.get(int(r['notes'][9:]), {}).get('english') != r['english']
                    or by_row.get(int(r['notes'][9:]), {}).get('kind') != 'text']
        check('every grey reference row points to a text row with the same English', not bad_refs,
              f'{len(refs)} reference rows' + (f'; bad: {bad_refs[:5]}' if bad_refs else ''))

        h1 = tree_hash(os.path.join(tmp, 'data/localized_data'))
        code, _ = run(imp, tmp)
        check('second import gives identical files', code == 0 and tree_hash(os.path.join(tmp, 'data/localized_data')) == h1)

        print('\n[4] Rebuild keeps translations and notes')
        code, out = run(build, tmp)
        wb, ws, col = load_sheet(os.path.join(tmp, 'translation_table.xlsx'))
        after = {r['key']: r for r in table_rows(ws, col) if r['kind'] == 'text'}
        lost = [k for k, c in filled.items() if after.get(k, {}).get('ukrainian') != c]
        lost_notes = [k for k, n in notes.items()
                      if ws.cell(row=after[k]['row'], column=col['translator_notes']).value != n]
        check('rebuild keeps every translation', code == 0 and not lost, f'lost: {lost[:3]}' if lost else f'{len(filled)} kept')
        check('rebuild keeps every translator note', not lost_notes)
        added = [k for k, r in after.items() if r['ukrainian'] and k not in filled]
        check('rebuild adds no translation nobody typed', not added, f'added: {added[:3]}' if added else '')

        print('\n[5] Gender change keeps translations')
        unknown = [r for r in after.values() if re.match(r'^[0-9a-f]+-S\?$', r['key']) and ',' not in r['speakers']
                   and genders.get(r['speakers']) == '?']
        if unknown:
            r = rnd.choice(unknown)
            ws.cell(row=r['row'], column=col['ukrainian']).value = fake(r['english'])
            wb.save(os.path.join(tmp, 'translation_table.xlsx'))
            from openpyxl import load_workbook
            gwb = load_workbook(os.path.join(tmp, 'speaker_genders.xlsx'))
            for gr in gwb['Genders'].iter_rows(min_row=2):
                if gr[0].value == r['speakers']:
                    gr[1].value = 'F'
            gwb.save(os.path.join(tmp, 'speaker_genders.xlsx'))
            code, out = run(build, tmp)
            wb, ws, col = load_sheet(os.path.join(tmp, 'translation_table.xlsx'))
            new = {x['key']: x for x in table_rows(ws, col) if x['kind'] == 'text'}
            moved = new.get(r['key'][:-1] + 'F')
            check(f"speaker {r['speakers']} ? -> F: translation moved to the female row",
                  code == 0 and moved and moved['ukrainian'] == fake(r['english']) and 'check gender forms' in moved['notes'])
            code, out = run(imp, tmp)
            check('import after the gender change', code == 0, re.findall(r'Consistency: .*', out))
        else:
            check('gender change test skipped (no single-speaker "?" row)', True)

        print('\n[6] Mistakes must be reported')
        wb, ws, col = load_sheet(os.path.join(tmp, 'translation_table.xlsx'))
        rows = table_rows(ws, col)
        one = rnd.choice([r for r in rows if r['kind'] == 'text' and r['lines'] == 1 and r['key'] not in filled])
        ref = rnd.choice([r for r in rows if r['kind'] == 'ref'])
        tagged = rnd.choice([r for r in rows if r['kind'] == 'text' and re.search(r'\{[A-Za-z]', r['english'])
                             and r['key'] not in filled])
        ws.cell(row=one['row'], column=col['ukrainian']).value = 'рядок один\nрядок два'
        ws.cell(row=ref['row'], column=col['ukrainian']).value = 'набрано в сірому рядку'
        ws.cell(row=tagged['row'], column=col['ukrainian']).value = re.sub(r'\{[^{}]*\}', '', fake(tagged['english']), count=1)
        wb.save(os.path.join(tmp, 'translation_table.xlsx'))
        code, out = run(imp, tmp)
        check('import fails (exit code 1)', code == 1)
        check(f"line break in one-line text reported (row #{one['row']})", f"table row #{one['row']}:" in out)
        check(f"typing in grey reference row reported (row #{ref['row']})", f"table row #{ref['row']} is a grey reference row" in out)
        check(f"deleted tag reported (row #{tagged['row']})", 'markup differs' in out)
        bad_tag_files = set(re.findall(r'== (\S+): \d+ error', out))
        check('only the file(s) with the deleted tag have errors', bad_tag_files and all(
            '{' in open(os.path.join(tmp, 'data/original_data', f), encoding='utf-8-sig').read() for f in bad_tag_files),
            f'{len(bad_tag_files)} file(s)')
    finally:
        if a.keep:
            print(f'\nKept: {tmp}')
        else:
            shutil.rmtree(tmp, ignore_errors=True)
    failed = [r for r in RESULTS if not r[1]]
    print(f"\n{'ALL PASSED' if not failed else f'{len(failed)} FAILED'}: {len(RESULTS) - len(failed)}/{len(RESULTS)} "
          f'checks, seed {seed}, {time.time() - t0:.0f}s')
    sys.exit(1 if failed else 0)


if __name__ == '__main__':
    main()
