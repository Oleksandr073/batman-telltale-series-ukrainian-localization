"""AI draft translation in batches: export untranslated rows with context, then write checked drafts into the table.

Usage (from the repo root):
  python ai_translate.py status TABLE
  python ai_translate.py export TABLE OUT_BATCH (--scene GAME_FILE | --next N) [--shown-only]
  python ai_translate.py apply  TABLE DRAFT_TSV [--replace-drafts]

  status   how many rows are translated: by people, AI drafts, empty - per episode/section.
  export   writes OUT_BATCH (text file) with the rows to translate and their context: who speaks to whom, gender
           columns, notes, the glossary names used in the batch, the ти/ви rules for the speaker pairs and the
           scene view to read (ai_context/scenes/...). Only rows with an EMPTY Ukrainian cell are exported.
             --scene     one game file, e.g. BM_pc_Batman101_data/env_mayorsoffice (rows first seen there)
             --next N    the next N empty rows in table (play) order
             --shown-only  only rows the official translations translated ("Official" = yes/some)
  apply    reads DRAFT_TSV (columns: key <TAB> ukrainian [<TAB> note for reviewers]; ¶ between lines), checks:
           same number of lines, markup copied in the same order, same leading spaces, same ending punctuation
           per line, no ASCII ' or " (use ’ and “ ”). Rows that pass are written into EMPTY Ukrainian cells only
           and marked "AI draft" in Translator notes; rows that fail are listed and not written.
             --replace-drafts  also replace cells that are still marked "AI draft" (never cells typed by people)

People review drafts in the table and delete "AI draft" from Translator notes when a row is checked.
Then the usual flow: npm run import, npm run review.
"""
import collections
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tm_lib import MARKUP, save_stable  # noqa: E402

DRAFT = 'AI draft'
GLOSSARY = 'ai_context/glossary.md'
SPEAKERS = 'ai_context/speakers.md'
SCENES = 'ai_context/scenes'
LATIN = re.compile(r'[A-Za-z]')
END = re.compile(r'(\.\.\.|…|--|[.?!,:;—–\-])?(["”’)\]]*)$')   # ending punctuation, then closing quotes/brackets


def load(path):
    from openpyxl import load_workbook
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from tm_lib import TABLE_COLUMNS
    wb = load_workbook(path)
    ws = wb['Translate']
    titles = [c.value for c in ws[1]]
    col = {name: titles.index(title) + 1 for name, title, _ in TABLE_COLUMNS if title in titles}
    return wb, ws, col


def rows_of(ws, col):
    """All rows in table order: dicts with the column values, plus 'row', 'kind' and 'section'."""
    out, section, scene, title = [], '', '', ''
    names = list(col)
    for vals in ws.iter_rows(min_row=2, values_only=True):
        r = {n: ('' if vals[col[n] - 1] is None else str(vals[col[n] - 1])) for n in names}
        r['row'] = len(out) + 2
        if r.get('key'):
            r['kind'] = 'text'
        elif r.get('notes', '').startswith('see row #'):
            r['kind'] = 'ref'
        elif r.get('english') and not r.get('first_seen'):
            r['kind'], section = 'section', r['english']
        else:
            r['kind'], scene, title = 'scene', r.get('first_seen', ''), r.get('english', '')
        r['section'], r['scene'], r['scene_title'] = section, scene, title
        out.append(r)
    return out


def is_draft(r):
    return r.get('translator_notes', '').startswith(DRAFT)


# --- context: glossary and ти/ви ------------------------------------------------------------------------------
def glossary():
    out = []
    for line in open(GLOSSARY, encoding='utf-8'):
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if line.startswith('|') and len(cells) >= 2 and cells[0] not in ('English', '') and not cells[0].startswith('-'):
            out.append((cells[0], cells[1], cells[3] if len(cells) > 3 else ''))
    return out


ALIASES = {'HARVEY': ['DENT', 'TWOFACE'], 'MONTOYA': ['RENEE'], 'OSWALD': ['COBBLEPOT'], 'VICKI': ['VALE'],
           'REPORTERS': ['REPORTER', 'RAMARQUE', 'SPENCER'], 'JOKER': ['JOHN', 'JOHNDOE']}


def ids_for(name):
    ids = set()
    for part in name.split('/'):
        words = re.sub(r'[^A-Za-z ]', '', part).upper().split()
        if words:
            ids.update({''.join(words), words[-1]})
    for k in list(ids):
        ids.update(ALIASES.get(k, []))
    return ids


def norm_id(s):
    s = re.sub(r'[^A-Z]', '', s.upper())
    return 'BRUCE' if s in ('CHOICEUI', 'BATMANBRUCE') else s


def tyvy_rules():
    rules = []
    for line in open(SPEAKERS, encoding='utf-8'):
        c = [x.strip() for x in line.strip().strip('|').split('|')]
        if len(c) >= 8 and '→' in c[0] and '→' in c[3]:
            for pair, dec in ((c[0], c[2]), (c[3], c[5])):
                a, b = [x.strip() for x in pair.split('→')]
                rules.append((ids_for(a), ids_for(b), f"{pair}: {dec.strip('*')}" + (f' ({c[6]})' if c[6] else '')))
    return rules


def match_id(sid, ids):
    return sid in ids or any(sid.startswith(i) for i in ids if i == 'REPORTER')


# --- export -------------------------------------------------------------------------------------------------
def export(table, out, scene=None, nxt=None, shown_only=False):
    _, ws, col = load(table)
    rows = rows_of(ws, col)
    todo = [r for r in rows if r['kind'] == 'text' and not r['ukrainian'].strip()]
    if shown_only:
        todo = [r for r in todo if r.get('official', '').startswith(('yes', 'some'))]
    if scene:
        s = scene.replace('\\', '/').removesuffix('.txt').removesuffix('_english')
        todo = [r for r in todo if r['first_seen'].split(':')[0] == s + '_english.txt']
    if nxt:
        todo = todo[:nxt]
    if not todo:
        sys.exit('Nothing to export: no empty Ukrainian cells match.')
    gl = glossary()
    text = ' '.join(r['english'] for r in todo)
    names = [g for g in gl if re.search(r'(?<![A-Za-z])' + re.escape(g[0]) + r'(?![a-z])', text)]
    rules, pairs = tyvy_rules(), collections.Counter()
    for r in todo:
        sp = [norm_id(s) for s in r['speakers'].split(', ') if s]
        li = [norm_id(s) for s in r['listeners'].split(', ') if s]
        if sp and li:
            for a, b, txt in rules:
                if match_id(sp[0], a) and match_id(li[0], b):
                    pairs[txt] += 1
    files = sorted({r['first_seen'].split(':')[0] for r in todo})
    L = [f'# Batch: {len(todo)} rows, {sum(len(r["english"].split()) for r in todo)} English words',
         f'# Table: {table}. Translate each row; write the draft as: key<TAB>ukrainian (¶ between lines).',
         '# Rules: ai_context/style_guide.md (read it), names below, ти/ви below; keep markup and ending punctuation.',
         '', '## Scenes to read for context (all lines, conversation order)']
    L += [f'- {SCENES}/{f}' for f in files]
    L += ['', '## Glossary names in this batch'] + [f'- {e} → {u}' + (f'  ({n})' if n else '') for e, u, n in names]
    L += ['', '## ти/ви for the speaker pairs in this batch (from speakers.md)']
    L += [f'- {t}  [{n} rows]' for t, n in pairs.most_common()] or ['- (no listed pair; see the general rules in speakers.md)']
    L += ['', '## Rows']
    last_scene = None
    for r in todo:
        if r['scene'] != last_scene:
            last_scene = r['scene']
            L += ['', f'=== {r["section"]} / {r["scene_title"]}  ({r["scene"]})']
        who = r['speakers'] or '?'
        to = f" → {r['listeners']} (guess)" if r['listeners'] else ''
        g = ', '.join(x for x in (f"speaker {r['speaker_gender']}" if r['speaker_gender'] else '',
                                  f"listener {r['listener_gender']}" if r['listener_gender'] else '') if x)
        L += ['', f"[{r['key']}] row {r['row']} | {who}{to}" + (f' | gender: {g}' if g else '')
              + f" | lines {r['lines']} | official {r.get('official', '')} | {r['first_seen']}"]
        if r['notes']:
            L.append(f"  notes: {r['notes']}")
        if r.get('translator_notes'):
            L.append(f"  translator notes: {r['translator_notes']}")
        L.append(f"  EN: {r['english']}")
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    with open(out, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(L) + '\n')
    print(f'{len(todo)} rows -> {out} (scenes: {", ".join(f.rsplit("/", 1)[-1] for f in files[:5])}'
          + (' ...' if len(files) > 5 else '') + ')')


# --- checks -------------------------------------------------------------------------------------------------
def split_cell(cell):
    return [p for p in re.split(r' ?¶ ?|\r\n|\r|\n', cell)]


def check_row(en, uk, lines):
    """-> list of problems (empty = ok)."""
    p = []
    e_parts, u_parts = en.split(' ¶ '), split_cell(uk)
    if len(u_parts) != int(lines or 1):
        return [f'{len(u_parts)} line(s), needs {lines}']
    if [m for m in MARKUP.findall(en)] != [m for m in MARKUP.findall(uk)]:
        p.append(f'markup differs: needs {MARKUP.findall(en)}, has {MARKUP.findall(uk)}')
    for i, (a, b) in enumerate(zip(e_parts, u_parts), 1):
        where = f'line {i}: ' if len(e_parts) > 1 else ''
        if len(a) - len(a.lstrip(' ')) != len(b) - len(b.lstrip(' ')):
            p.append(f'{where}leading spaces differ')
        ea, eb = END.search(a.rstrip()).groups(), END.search(b.rstrip()).groups()
        if ea != eb and a.strip() and b.strip():
            p.append(f'{where}ending {"".join(x or "" for x in ea)!r} -> {"".join(x or "" for x in eb)!r}')
        if a.count('--') != b.count('--') or a.count('...') != b.count('...') or a.count('…') != b.count('…'):
            p.append(f'{where}"--" / "..." / "…" count differs')
    vis = MARKUP.sub('', uk)
    for ch, name in (("'", "ASCII ' (use ’)"), ('"', 'ASCII " (use “ ”)')):
        if ch in vis and ch not in MARKUP.sub('', en):
            p.append(name)
    return p


def latin_left(en, uk):
    """Latin words in the draft that are not in the English markup - a warning (names like GCPD may be fine)."""
    return re.findall(r'[A-Za-z]{2,}', re.sub(r'\[[^\]]*\]', '', MARKUP.sub('', uk)))


# --- apply --------------------------------------------------------------------------------------------------
def apply(table, draft, replace_drafts=False):
    wb, ws, col = load(table)
    rows = {r['key']: r for r in rows_of(ws, col) if r['kind'] == 'text'}
    ok, bad, skipped, warn = 0, [], [], []
    seen = set()
    with open(draft, encoding='utf-8-sig') as fh:
        for n, line in enumerate(fh, 1):
            line = line.rstrip('\r\n')
            if not line.strip() or line.startswith('#'):
                continue
            key, _, uk = line.partition('\t')
            uk, _, note = uk.partition('\t')
            key, uk, note = key.strip().strip('[]'), uk.strip(' '), note.strip()
            r = rows.get(key)
            if not r:
                bad.append(f'{draft}:{n}: key {key!r} not in the table')
                continue
            if key in seen:
                bad.append(f'{draft}:{n}: key {key} twice in the draft')
                continue
            seen.add(key)
            if not uk.strip():
                continue
            if r['ukrainian'].strip() and not (replace_drafts and is_draft(r)):
                skipped.append(f"row {r['row']}: already has a translation" + (' (AI draft)' if is_draft(r) else ''))
                continue
            probs = check_row(r['english'], uk, r['lines'])
            if probs:
                bad.append(f"row {r['row']} [{key}]: " + '; '.join(probs) + f"\n      EN: {r['english']}\n      UK: {uk}")
                continue
            lat = latin_left(r['english'], uk)
            if lat:
                warn.append(f"row {r['row']}: Latin letters {lat[:5]} - fine only for names that stay Latin")
            ws.cell(row=r['row'], column=col['ukrainian']).value = uk
            notes = r.get('translator_notes', '')
            if notes.startswith(DRAFT):           # replacing a draft: drop its old mark and note
                notes = notes.split('; ', 1)[1] if '; ' in notes else ''
            mark = DRAFT + (f': {note}' if note else '')
            ws.cell(row=r['row'], column=col['translator_notes']).value = mark + (f'; {notes}' if notes else '')
            ok += 1
    if ok:
        try:
            save_stable(wb, table)
        except PermissionError:
            sys.exit(f'Cannot write {table} - close it in Excel / LibreOffice first. Nothing was written.')
    print(f'Wrote {ok} AI draft(s) into {table}; {len(bad)} row(s) rejected; {len(skipped)} skipped (already filled).')
    for b in bad:
        print('  REJECTED ' + b)
    for w in warn:
        print('  WARNING  ' + w)
    for s in skipped[:10]:
        print('  SKIPPED  ' + s)
    return 1 if bad else 0


def status(table):
    _, ws, col = load(table)
    c = collections.defaultdict(collections.Counter)
    for r in rows_of(ws, col):
        if r['kind'] != 'text':
            continue
        st = 'empty' if not r['ukrainian'].strip() else 'AI draft' if is_draft(r) else 'by people'
        for sec in (r['section'], 'ALL'):
            c[sec][st] += 1
            if r.get('official', '').startswith(('yes', 'some')):
                c[sec][st + ' (shown)'] += 1
    print(f"{'section':60s} {'people':>8} {'AI draft':>9} {'empty':>8}   (shown in game: people/draft/empty)")
    for sec, v in c.items():
        print(f"{sec[:60]:60s} {v['by people']:8d} {v['AI draft']:9d} {v['empty']:8d}   "
              f"{v['by people (shown)']}/{v['AI draft (shown)']}/{v['empty (shown)']}")


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    a = sys.argv[1:]
    flags = {f: (a[a.index(f) + 1] if f in ('--scene', '--next') else True) for f in
             ('--scene', '--next', '--shown-only', '--replace-drafts') if f in a}
    pos = [x for i, x in enumerate(a) if not x.startswith('--') and not (i and a[i - 1] in ('--scene', '--next'))]
    if pos[:1] == ['status'] and len(pos) == 2:
        return status(pos[1])
    if pos[:1] == ['export'] and len(pos) == 3 and ('--scene' in flags or '--next' in flags):
        return export(pos[1], pos[2], flags.get('--scene'), int(flags['--next']) if '--next' in flags else None,
                      '--shown-only' in flags)
    if pos[:1] == ['apply'] and len(pos) == 3:
        return apply(pos[1], pos[2], '--replace-drafts' in flags)
    sys.exit(__doc__)


if __name__ == '__main__':
    sys.exit(main())
